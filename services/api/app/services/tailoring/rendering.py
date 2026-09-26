"""Turning a validated document into a file the user can send.

DOCX rather than PDF, for one reason that outweighs the rest: a job applicant
almost always needs to edit the resume before sending it, and an applicant
tracking system almost always parses DOCX better than it parses a PDF laid out
by a generator. A PDF renderer can be added beside this one — the artifact row
records its own `media_type` and `filename` precisely so that a second format is
a new function here rather than a schema change.

Nothing in this module makes a decision about CONTENT. It receives a document
that has already passed grounding and lays it out. If a rule about what may
appear on a resume is ever needed, it belongs in `generation.py` (what to write)
or `grounding.py` (what may be stood behind), never here — a renderer that
filters is a second, invisible policy.

The bytes never touch PostgreSQL. `render` returns them and the caller hands
them to `app/services/storage.py`, which is the only module that knows where an
object lives.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.schemas.tailored_resume import TailoredResumeDocument, TailoredLink

#: What the Open XML word-processing format is called on the wire. Sent as the
#: artifact's Content-Type and stored on the row, so a future PDF renderer does
#: not need anything downstream to learn a second content type.
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

#: Anything outside this is replaced in a generated filename. Deliberately
#: strict: the filename is built from a job title and a company name, both of
#: which came from a third-party posting, and it ends up in a Content-Disposition
#: header. A quote or a newline there is a header-injection bug.
_UNSAFE_FILENAME = re.compile(r"[^A-Za-z0-9._-]+")

#: A run of dots, collapsed to one. Without this, "../../etc/passwd" sanitises
#: to "..-..-etc-passwd" — every separator gone and every ".." intact. Nothing
#: currently joins this name to a path (the storage key is derived from the
#: content hash, and this string only ever reaches a Content-Disposition
#: header), so today that is untidy rather than exploitable. It is collapsed
#: anyway: the value is attacker-controlled, and the next reader of this
#: function should not have to reconstruct that argument before deciding
#: whether they may use it as a path segment.
_DOT_RUN = re.compile(r"\.{2,}")

_MAX_FILENAME_STEM = 80

#: The gap between the contact header and the first section, carried by the
#: header's last paragraph rather than by an empty paragraph of its own.
HEADER_GAP = Pt(12)


def build_filename(*, job_title: str, company_name: str | None, version_number: int) -> str:
    """A name the user will recognise in their downloads folder.

    Includes the profile version number because the same job tailored from two
    versions produces two genuinely different documents, and a user who cannot
    tell them apart in a file listing will send the wrong one.
    """
    parts = [part for part in (company_name, job_title) if part]
    stem = _UNSAFE_FILENAME.sub("-", "-".join(parts))
    stem = _DOT_RUN.sub(".", stem).strip("-.") or "resume"
    return f"{stem[:_MAX_FILENAME_STEM]}-v{version_number}.docx"


def _format_date(value: date | None, precision: str | None) -> str:
    """One date, written only as precisely as it is known.

    The same discipline `formatDateWithPrecision` applies in the web app: the
    domain stores "May 2024" as 2024-05-01 with precision "month", and printing
    "1 May 2024" on a resume would put back exactly the false precision the
    precision column exists to prevent.
    """
    if value is None:
        return ""
    if precision == "year":
        return value.strftime("%Y")
    if precision == "day":
        return value.strftime("%d %b %Y")
    return value.strftime("%b %Y")


def _format_period(
    start: date | None,
    start_precision: str | None,
    end: date | None,
    end_precision: str | None,
) -> str:
    """A JOB's date range, where a missing end means the job is current.

    That reading belongs to experience alone: `Experience.end_date` is NULL for
    a role the candidate still holds, so "Mar 2022 – Present" is what the row
    says.
    """
    left = _format_date(start, start_precision)
    right = _format_date(end, end_precision) if end else "Present"
    if not left:
        return right if end else ""
    return f"{left} – {right}"


def _format_dates(
    start: date | None,
    start_precision: str | None,
    end: date | None,
    end_precision: str | None,
) -> str:
    """An education entry's dates, where a missing end means only that none was given.

    Not `_format_period`, and the difference is a fact about the candidate. The
    previous renderer shared the job reading, so a degree dated only "2018" —
    stored as a start date, because that is the one the resume reader found —
    printed as "(2018 – Present)": a claim of current enrolment nobody made. A
    degree with only an end date lost its date entirely, because the line was
    printed only when a start existed.

    A single date is printed as itself, at its own precision. Both are printed
    as a range, exactly as before.
    """
    left = _format_date(start, start_precision)
    right = _format_date(end, end_precision)
    if left and right:
        return f"{left} – {right}"
    return left or right


#: Characters XML 1.0 cannot carry at all: the C0 controls other than tab,
#: newline and carriage return, lone surrogates, and the two non-characters.
#: lxml refuses to write a string holding one ("All strings must be XML
#: compatible"), so a single form feed left in a bullet by PDF extraction —
#: which grounding rightly accepts, because the same byte is in the profile —
#: used to crash the render stage into a permanent `render_failed`.
_XML_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")

#: The two of those that are whitespace in every sense but XML's. Replaced by a
#: space rather than deleted, so "payments\x0cservice" does not print as
#: "paymentsservice".
_XML_ILLEGAL_SPACES = re.compile("[\x0b\x0c]")


def _clean(text: str | None) -> str:
    """Text the file format can hold. Removes only what carries no meaning.

    Not a content decision: every character removed here is one no reader could
    ever have seen, because no document format can display it. Everything a
    reader can see is left exactly as grounding approved it.
    """
    if not text:
        return ""
    return _XML_ILLEGAL.sub("", _XML_ILLEGAL_SPACES.sub(" ", text))


def _heading(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(_clean(text).upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


# --- contact links -------------------------------------------------------------------

#: A scheme at the front of a stored url: "https:", "mailto:", "tel:".
_SCHEME = re.compile(r"^([A-Za-z][A-Za-z0-9+.-]*):(?://)?")

#: The schemes a link on a resume is allowed to OPEN. The url is the candidate's
#: own, but the file goes to a stranger's machine, and a `file:` or UNC target
#: in a hyperlink is a thing Word will offer to open for them. Anything else is
#: still printed — the candidate wrote it — just not made clickable.
HYPERLINK_SCHEMES = frozenset({"http", "https", "mailto", "tel"})

#: A bare "github.com/alex" as a resume often writes it: a host and a path, no
#: scheme, no spaces.
_BARE_HOST = re.compile(r"^[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)+(?:[/:?#]\S*)?$")


def display_url(url: str) -> str:
    """What a link prints as: the address, without the parts nobody types.

    "https://www.linkedin.com/in/alex/" prints "linkedin.com/in/alex". The
    header used to print the LABEL instead — "LinkedIn | GitHub | Portfolio" —
    which on paper, and to an applicant tracking system reading the text layer,
    is three words with no address behind them. The address is the useful half.
    """
    shown = url.strip()
    scheme = _SCHEME.match(shown)
    # Only a scheme a link may open is dropped. A bare host with a port
    # ("alexmorgan.dev:8443/cv") looks like "scheme:" to the pattern, and
    # stripping it printed "8443/cv"; an unopenable scheme ("javascript:",
    # "skype:") is printed as written, because the candidate wrote it.
    if scheme and not _BARE_HOST.match(shown) and scheme.group(1).casefold() in HYPERLINK_SCHEMES:
        shown = shown[scheme.end():]
    if shown[:4].casefold() == "www.":
        shown = shown[4:]
    return shown.rstrip("/")


def hyperlink_target(url: str) -> str | None:
    """Where a link should open, or None when it must not be made clickable.

    A stored url with a permitted scheme is used exactly as stored. A bare host
    ("github.com/alex") is given "https://", because Word treats a target with
    no scheme as a path on the reader's own disk. Anything else is None.
    """
    url = url.strip()
    match = _SCHEME.match(url)
    if match and not _BARE_HOST.match(url):
        return url if _openable(url) else None
    if _BARE_HOST.match(url):
        return f"https://{url}"
    return None


def _openable(target: str | None) -> bool:
    """A url with a scheme a resume link may open, and something after it."""
    match = _SCHEME.match(target or "")
    return (
        match is not None
        and match.group(1).casefold() in HYPERLINK_SCHEMES
        and len(target or "") > len(match.group(0))
    )


def _add_link(paragraph, link: TailoredLink) -> None:
    """One contact link: a real hyperlink, printed as its address.

    A `w:hyperlink` whose `r:id` names an EXTERNAL relationship to the full url
    — the shape Word, LibreOffice and a parser reading relationships all open.
    python-docx has no call for writing one, so the element is built here around
    an ordinary run, which keeps the run's formatting API.
    """
    text = _clean(display_url(link.url)) or _clean(link.label) or _clean(link.url)
    run = paragraph.add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # A url carrying a character XML cannot hold is not a working address, and
    # writing it into the relationship crashed the render — a permanent
    # render_failed for a document grounding had approved. Printed, not linked.
    target = hyperlink_target(link.url) if _clean(link.url) == link.url else None
    if target is None:
        return
    run.font.underline = True
    relationship_id = paragraph.part.relate_to(
        target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.set(qn("w:history"), "1")
    run._r.addprevious(hyperlink)
    hyperlink.append(run._r)


def _set_properties(document, tailored: TailoredResumeDocument) -> None:
    """The file's own metadata: the candidate's name, and nothing of ours.

    python-docx's template ships `creator = "python-docx"` and a description of
    "generated by python-docx", and every tailored resume carried both into the
    employer's File > Properties. That is the same exhaust as a "Generated by
    EasyApply" footer, one dialog box deeper. The author is the person whose
    resume it is — the name the document already prints at the top — or blank.
    """
    name = " ".join(_clean(tailored.contact.full_name).split()) if tailored.contact else ""
    properties = document.core_properties
    properties.author = name
    properties.last_modified_by = name
    # `comments` is dc:description, the field the template's sentence lives in.
    for field in ("title", "subject", "comments", "keywords", "category"):
        setattr(properties, field, "")
    # The template's created/modified stamps (2013-12-23T23:15:00Z, a known
    # python-docx fingerprint) told an employer the resume was written in 2013.
    # Removed rather than set to "now", which would make two renders of one
    # document differ.
    element = properties._element
    for tag in ("dcterms:created", "dcterms:modified"):
        for child in element.findall(qn(tag)):
            element.remove(child)


def render(tailored: TailoredResumeDocument) -> bytes:
    """Lay out a tailored resume as a DOCX file.

    Takes the document and nothing else. It used to take the job title and the
    company too, and print them at the top as "Prepared for X at Y" — which put
    the posting's words, and the fact that this file was aimed at that vacancy,
    on a page the candidate sends to that employer. Every word in the output now
    comes from the tailored document, and the tailored document comes from the
    candidate's own confirmed profile.

    A line with nothing printable in it is not given a paragraph. It is not a
    claim — there is nothing in it to claim — and an empty paragraph is a hole
    in the layout that `verify` refuses.
    """
    document = Document()
    _set_properties(document, tailored)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    contact = tailored.contact
    # One line, whatever the stored value holds: a name read as "Ann\r\nLee"
    # became a line break in the file and failed `verify`'s name check.
    full_name = " ".join(_clean(contact.full_name).split()) if contact is not None else ""
    # The last paragraph of the header block carries the gap between the header
    # and the first section. That gap used to be an empty spacer paragraph in
    # every document — which an ATS reads as a blank line and a person editing
    # the file finds as a paragraph they cannot see.
    header_end = None

    # The name, at the top, at the size a name goes. Everything below it is
    # about a career; this line is the only thing that says whose.
    if full_name.strip():
        name_paragraph = document.add_paragraph()
        name_paragraph.paragraph_format.space_after = Pt(0)
        run = name_paragraph.add_run(full_name)
        run.bold = True
        run.font.size = Pt(18)
        header_end = name_paragraph

    headline = _clean(tailored.headline.text) if tailored.headline is not None else ""
    if headline.strip():
        title = document.add_paragraph()
        title.paragraph_format.space_after = Pt(0)
        run = title.add_run(headline)
        # Under a name this is a professional headline. With no name above it —
        # a version holding no contact columns — it is all the document has, so
        # it takes the size the name would have had.
        run.bold = True
        run.font.size = Pt(11 if full_name.strip() else 16)
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        header_end = title

    if contact is not None:
        details = [
            _clean(value)
            for value in (contact.location, contact.phone, contact.email)
            if _clean(value).strip()
        ]
        if details:
            line = document.add_paragraph()
            line.paragraph_format.space_after = Pt(0)
            run = line.add_run("  |  ".join(details))
            run.font.size = Pt(9.5)
            header_end = line

        links = [link for link in contact.links if _clean(link.url).strip()]
        if links:
            line = document.add_paragraph()
            line.paragraph_format.space_after = Pt(0)
            for position, link in enumerate(links):
                if position:
                    separator = line.add_run("  |  ")
                    separator.font.size = Pt(9.5)
                    separator.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                _add_link(line, link)
            header_end = line

    if header_end is not None:
        header_end.paragraph_format.space_after = HEADER_GAP

    summary = [_clean(line.text) for line in tailored.summary]
    summary = [text for text in summary if text.strip()]
    if summary:
        _heading(document, "Summary")
        for text in summary:
            # A paragraph, not a bullet. The summary is the one place the
            # document speaks in prose, and bulleting a single sentence makes it
            # read as the first item of a list the reader then goes looking for.
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.space_after = Pt(0)

    groups = [
        (_clean(group.category), [_clean(name) for name in group.skills if _clean(name).strip()])
        for group in tailored.skills
    ]
    groups = [(category, names) for category, names in groups if names]
    if groups:
        _heading(document, "Skills")
        for category, names in groups:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            label = paragraph.add_run(f"{category}: ")
            label.bold = True
            paragraph.add_run(", ".join(names))

    if tailored.experience:
        _heading(document, "Experience")
        for entry in tailored.experience:
            header = document.add_paragraph()
            header.paragraph_format.space_after = Pt(0)
            role = header.add_run(_clean(f"{entry.job_title}, {entry.company}"))
            role.bold = True

            meta_text = _format_period(
                entry.start_date,
                entry.start_date_precision,
                entry.end_date,
                entry.end_date_precision,
            )
            location = _clean(entry.location)
            if location.strip():
                meta_text = f"{meta_text} · {location}" if meta_text else location
            if meta_text:
                meta = document.add_paragraph()
                meta.paragraph_format.space_after = Pt(2)
                run = meta.add_run(meta_text)
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

            _bullets(document, entry.bullets)

    if tailored.projects:
        _heading(document, "Projects")
        for project in tailored.projects:
            # A header is written only with something in it: a blank name and
            # no role used to give an empty paragraph, which `verify` rightly
            # rejects — a permanent render_failed for a grounded document.
            project_name = _clean(project.name).strip()
            role = _clean(project.role).strip()
            if project_name or role:
                header = document.add_paragraph()
                header.paragraph_format.space_after = Pt(0)
                if project_name:
                    name = header.add_run(project_name)
                    name.bold = True
                if role:
                    header.add_run(f" — {role}" if project_name else role)

            technologies = [_clean(name) for name in project.technologies if _clean(name).strip()]
            if technologies:
                line = document.add_paragraph()
                line.paragraph_format.space_after = Pt(2)
                run = line.add_run(", ".join(technologies))
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

            _bullets(document, project.bullets)

    if tailored.education:
        _heading(document, "Education")
        for entry in tailored.education:
            institution_text = _clean(entry.institution).strip()
            detail = ", ".join(
                _clean(part).strip()
                for part in (entry.degree, entry.field_of_study)
                if _clean(part).strip()
            )
            period = _format_dates(
                entry.start_date,
                entry.start_date_precision,
                entry.end_date,
                entry.end_date_precision,
            )
            if not (institution_text or detail or period):
                continue
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            if institution_text:
                institution = paragraph.add_run(institution_text)
                institution.bold = True
            if detail:
                paragraph.add_run(f" — {detail}" if institution_text else detail)
            if period:
                paragraph.add_run(f"  ({period})" if (institution_text or detail) else period)

    if tailored.certifications:
        _heading(document, "Certifications")
        for entry in tailored.certifications:
            certification_name = _clean(entry.name).strip()
            issuer = _clean(entry.issuing_organization).strip()
            if not (certification_name or issuer or entry.issue_date):
                continue
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            if certification_name:
                name = paragraph.add_run(certification_name)
                name.bold = True
            if issuer:
                paragraph.add_run(f" — {issuer}" if certification_name else issuer)
            if entry.issue_date:
                # The year, not the month. `Certification` has no precision
                # column, so a stored 2023-01-01 may be "January 2023" or may be
                # a resume that said "2023" — and there is no way to tell from
                # here. Printing the month would invent that distinction half
                # the time; printing the year is never wrong, only less. One
                # date, never a range: a certification has no end for
                # "Present" to stand in for.
                paragraph.add_run(f"  ({entry.issue_date.strftime('%Y')})")

    # No footer, and no "prepared for this role" line at the top either. Both
    # used to be here and both were a mistake of the same kind: this file is
    # sent to an employer under the candidate's name, and a line telling that
    # employer which tool produced it — or that a machine aimed it at their
    # vacancy — is our exhaust printed on the user's resume. Where a document
    # came from belongs on the review screen, which is ours.

    buffer = io.BytesIO()
    document.save(buffer)
    return _with_fixed_timestamps(buffer.getvalue())


#: The earliest time a zip entry can carry. python-docx stamps every part with
#: the moment it was written (2-second resolution), so two renders of one
#: document differed whenever they straddled a tick — and the stored artifact is
#: keyed by its SHA-256.
_ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)


def _with_fixed_timestamps(data: bytes) -> bytes:
    """The same package, byte-for-byte reproducible: each part's content,
    order, compression and attributes kept, only its timestamp fixed."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=_ZIP_EPOCH)
            fixed.compress_type = info.compress_type
            fixed.external_attr = info.external_attr
            target.writestr(fixed, source.read(info))
    return out.getvalue()


def _bullets(document, lines) -> None:
    for line in lines:
        text = _clean(line.text)
        if not text.strip():
            continue
        paragraph = document.add_paragraph(text, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(0)


# --- checking the file that was written ----------------------------------------------


class RenderVerificationError(ValueError):
    """The bytes `render` produced are not a file we will hand over.

    Carries a short machine-readable reason for the log. Never document text:
    the reason is logged, and the document is the user's career history.
    """

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


def verify(data: bytes, tailored: TailoredResumeDocument) -> None:
    """Re-open the rendered file and check it is the document we meant to write.

    Grounding approved the DOCUMENT. Nothing had ever looked at the FILE: the
    ten documents the validation pass downloaded carried no hyperlinks at all,
    and nothing in production would have noticed a file that opened empty, lost
    the name, or gained a blank paragraph. This is the last look before the
    bytes are stored, and it reads them the way a recipient's tooling would —
    with a DOCX reader, not by trusting the code that wrote them.

      * it opens as a word-processing document with paragraphs in it;
      * the candidate's name, when the document has one, is in its text;
      * no paragraph is empty — `render` writes none on purpose, so one here
        is a layout defect;
      * every `w:hyperlink` resolves to an EXTERNAL hyperlink relationship whose
        target is a url a resume link may open, and every hyperlink
        relationship the part holds is such a url.

    Raises `RenderVerificationError`; the pipeline turns that into
    `render_failed`, with no artifact.
    """
    try:
        reopened = Document(io.BytesIO(data))
    except Exception as exc:  # a zip that is not a document, or not a zip at all
        raise RenderVerificationError("does_not_open") from exc

    paragraphs = reopened.paragraphs
    if not paragraphs:
        raise RenderVerificationError("no_paragraphs")
    if any(not paragraph.text.strip() for paragraph in paragraphs):
        raise RenderVerificationError("empty_paragraph")

    name = " ".join(_clean(tailored.contact.full_name).split()) if tailored.contact else ""
    if name.strip() and name not in "\n".join(paragraph.text for paragraph in paragraphs):
        raise RenderVerificationError("name_missing")

    part = reopened.part
    for relationship in part.rels.values():
        if relationship.reltype != RELATIONSHIP_TYPE.HYPERLINK:
            continue
        if not relationship.is_external or not _openable(relationship.target_ref):
            raise RenderVerificationError("hyperlink_not_external")

    for element in part.element.body.iter(qn("w:hyperlink")):
        relationship_id = element.get(qn("r:id"))
        relationship = part.rels.get(relationship_id) if relationship_id else None
        if (
            relationship is None
            or relationship.reltype != RELATIONSHIP_TYPE.HYPERLINK
            or not relationship.is_external
            or not _openable(relationship.target_ref)
        ):
            raise RenderVerificationError("hyperlink_unresolved")
