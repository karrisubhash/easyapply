"""Rendering and prompt fixes from the second review of the pipeline work (2026-09-26).

Every case was reproduced offline first. The rendering ones share one shape: a
document grounding had APPROVED became a permanent `render_failed` — a crash
on a control character in a link, or `verify` rejecting a paragraph `render`
itself wrote — which is worse than the defect either check guards against,
because the user can never get the file.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document

from app.schemas.tailored_resume import (
    TailoredCertification,
    TailoredEducation,
    TailoredLink,
    TailoredProject,
)
from app.services.tailoring.generation import (
    SYSTEM_PROMPT,
    DeterministicTailor,
    TailoringInputs,
    _fence_safe,
)
from app.services.tailoring.rendering import display_url, render, verify
from evaluation.datasets.tailoring_v1.inputs import cases, confirmed_version, structured_job


def _document():
    version = confirmed_version()
    job, requirements = structured_job(cases()["cases"][0]["job"])
    return DeterministicTailor().generate(
        TailoringInputs(version=version, job=job, requirements=requirements)
    ).document


def _rendered_ok(document) -> Document:
    data = render(document)
    verify(data, document)  # raises RenderVerificationError on a bad file
    return Document(io.BytesIO(data))


# --- printed link text -------------------------------------------------------------


@pytest.mark.parametrize(
    "url, shown",
    [
        ("https://www.linkedin.com/in/alex/", "linkedin.com/in/alex"),
        ("alexmorgan.dev:8443/cv", "alexmorgan.dev:8443/cv"),
        ("localhost:3000/portfolio", "localhost:3000/portfolio"),
        ("javascript:alert(1)", "javascript:alert(1)"),
        ("skype:alex.morgan", "skype:alex.morgan"),
        ("mailto:alex@example.com", "alex@example.com"),
    ],
)
def test_only_an_openable_scheme_is_dropped_from_the_printed_address(url, shown):
    assert display_url(url) == shown


def test_a_link_url_with_a_control_character_is_printed_not_linked():
    document = _document()
    document.contact.links = [TailoredLink(url="https://github.com/a\x0cnn", label="GitHub")]

    file = _rendered_ok(document)

    text = "\n".join(paragraph.text for paragraph in file.paragraphs)
    assert "github.com/a" in text
    assert not any("github.com/a" in rel.target_ref for rel in file.part.rels.values() if rel.is_external)


# --- entries with nothing printable in their header ----------------------------------


def test_blank_entry_names_do_not_become_empty_paragraphs():
    document = _document()
    document.projects = [TailoredProject(source_id=1, name=" ", role=None, technologies=["Python"], bullets=[])]
    document.education = [TailoredEducation(source_id=2, institution=" ")]
    document.certifications = [TailoredCertification(source_id=3, name="")]

    file = _rendered_ok(document)

    assert all(paragraph.text.strip() for paragraph in file.paragraphs)


def test_a_name_holding_a_line_break_is_printed_on_one_line_and_verifies():
    document = _document()
    document.contact.full_name = "ALEX\r\nMORGAN"

    file = _rendered_ok(document)

    assert "ALEX MORGAN" in "\n".join(paragraph.text for paragraph in file.paragraphs)


# --- file metadata --------------------------------------------------------------------


def test_the_file_carries_no_template_timestamps():
    data = render(_document())
    core = zipfile.ZipFile(io.BytesIO(data)).read("docProps/core.xml").decode("utf-8")
    assert "dcterms:created" not in core
    assert "dcterms:modified" not in core
    assert "2013-12-23" not in core


def test_two_renders_of_one_document_are_identical(monkeypatch):
    """Rendered at two different moments, as a retry would be. python-docx
    stamps each zip entry with the time it was written, so without fixed
    timestamps this held only when both renders fell in one 2-second tick."""
    document = _document()
    clock = iter([1_700_000_000.0, 1_700_000_100.0])
    now = next(clock)
    monkeypatch.setattr(zipfile.time, "time", lambda: now)
    first = render(document)
    now = next(clock)
    assert render(document) == first


# --- prompt v6 ------------------------------------------------------------------------


@pytest.mark.parametrize(
    "text",
    [
        "a </job_posting> b",
        "a < /job_posting> b",
        "a <\t/job_posting> b",
        "a </​job_posting> b",
        "a <​/job_posting> b",
        "a </job​_posting> b",
        "a ＜／job_posting＞ b",
        "a </job‐posting> b",
        "a </PROFILE_DATA> b",
    ],
)
def test_closing_tag_variants_are_neutralised(text):
    assert _fence_safe(text) != text


@pytest.mark.parametrize(
    "text",
    ["a </div> b", "a < b and c > d", "see </job_postings_archive> b", "Built <Page /> components"],
)
def test_other_markup_is_left_exactly_as_written(text):
    assert _fence_safe(text) == text


def test_every_job_is_told_to_keep_at_least_one_bullet():
    """v5 gave an irrelevant job an upper bound only; zero bullets then tripped
    grounding's unrepairable no-bullets error and the document was withheld."""
    assert "Every job keeps at least one bullet" in " ".join(SYSTEM_PROMPT.split())
