"""The tailoring pipeline: generate, validate, render — and where it stops.

The state machine matters here for one reason above the others: a document that
fails grounding must reach `completed` without an artifact, not `failed`. Those
are different facts about different things ("we produced this and will not vouch
for it" versus "we could not produce anything") and a UI that cannot tell them
apart cannot tell the user what to do next.
"""

from __future__ import annotations

from datetime import date

import pytest

from app.models.tailoring import (
    GROUNDING_FAILED,
    GROUNDING_PASSED,
    SEVERITY_ERROR,
    STATUS_COMPLETED,
    STATUS_FAILED,
    TailoringRun,
)
from app.schemas.tailored_resume import (
    TailoredExperience,
    TailoredLine,
    TailoredResumeDocument,
)
from app.services.storage import InMemoryStorage
from app.services.tailoring import pipeline, service
from app.services.tailoring.generation import GenerationResult
from app.services.tailoring.providers import ProviderError
from tests.matching_fixtures import persisted_profile_version
from tests.tailoring_fixtures import a_backend_job, a_tailorable_version


@pytest.fixture
def storage():
    return InMemoryStorage()


@pytest.fixture
def setup(db, user):
    version = a_tailorable_version(db, user)
    job, requirements = a_backend_job(db, user)
    return version, job, requirements


def a_run(db, user, version, job, requirements, **overrides) -> TailoringRun:
    run = TailoringRun(
        user_id=user.id,
        job_id=job.id,
        profile_version_id=version.id,
        job_requirements_id=requirements.id,
        attempt=overrides.pop("attempt", 1),
        generator=overrides.pop("generator", "deterministic"),
        **overrides,
    )
    db.add(run)
    db.commit()
    db.refresh(run)
    return run


# --- the happy path ---------------------------------------------------------------


def test_a_run_completes_and_produces_an_approved_artifact(db, user, setup, storage):
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert run.status == STATUS_COMPLETED
    assert run.finished_at is not None
    assert run.failed_stage is None

    tailored = run.result
    assert tailored is not None
    assert tailored.grounding_status == GROUNDING_PASSED
    assert tailored.error_count == 0
    assert tailored.artifact is not None
    assert tailored.artifact.byte_size > 0
    # DOCX is a zip. If the renderer produced anything else, this is where it
    # shows up rather than in a user's download folder.
    assert storage.get(tailored.artifact.storage_key).startswith(b"PK")


def test_the_deterministic_generator_only_quotes_the_profile(db, user, setup, storage):
    """The property the default configuration rests on.

    Not "usually grounded" and not "grounded in practice": every bullet's text
    is byte-identical to the snapshot text it cites, so there is nothing for the
    validator to object to. If this ever fails, the default install has started
    writing prose.
    """
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    document = TailoredResumeDocument.model_validate(run.result.payload)
    lines = [
        line
        for entry in document.experience
        for line in entry.bullets
    ] + list(document.summary)

    assert lines, "the fixture profile has prose; the generator selected none of it"
    assert all(line.text == line.source_quote for line in lines)


def test_selection_leads_with_what_the_posting_asked_for(db, user, setup, storage):
    """Tailoring's actual product: order, not invention."""
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    document = TailoredResumeDocument.model_validate(run.result.payload)
    senior = next(e for e in document.experience if e.company == "Acme Corp")

    assert "Python" in senior.bullets[0].text
    # The office move is true and irrelevant. It must not lead.
    assert "office move" not in senior.bullets[0].text


def test_the_document_records_which_requirements_it_aimed_at(db, user, setup, storage):
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    document = TailoredResumeDocument.model_validate(run.result.payload)
    assert "Python" in document.targeted_requirements


# --- grounding failure is a verdict, not a crash -------------------------------------


def test_a_document_that_fails_grounding_completes_without_an_artifact(
    db, user, setup, storage, monkeypatch
):
    """The single most important behaviour in this domain.

    A generator that invents a number produces a COMPLETED run carrying the
    document, the verdict and the findings — and no file. The user can see
    exactly what was written and exactly why it is not being handed over.

    Driven as a DETERMINISTIC run, which is where this contract is unconditional.
    An AI run reaches the reconciliation pass first: one invented number in one
    bullet is removed and the rest of the document is handed over, which is the
    whole point of that pass. The version of this contract that survives it —
    an AI document with nothing groundable left in it is still withheld — is
    `test_tailoring_ai.py::test_a_model_document_that_cannot_be_rescued_is_still_withheld`.
    """
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements, generator="deterministic")

    experience = version.experiences[0]

    class InventingTailor:
        name = "deterministic"
        model = None

        def generate(self, inputs):
            return GenerationResult(
                document=TailoredResumeDocument(
                    experience=[
                        TailoredExperience(
                            source_id=experience.id,
                            company=experience.company,
                            job_title=experience.job_title,
                            location=experience.location,
                            start_date=experience.start_date,
                            start_date_precision=experience.start_date_precision,
                            end_date=experience.end_date,
                            end_date_precision=experience.end_date_precision,
                            bullets=[
                                TailoredLine(
                                    text="Ran the team's on-call rotation for 40 engineers.",
                                    source_quote="Ran the team's on-call rotation.",
                                    source_field="description",
                                )
                            ],
                        )
                    ]
                ),
                generator="deterministic",
                model=None,
                prompt_version="test",
            )

    monkeypatch.setattr(pipeline, "get_tailor", lambda _generator=None: InventingTailor())

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert run.status == STATUS_COMPLETED
    assert run.failed_stage is None
    assert run.error_code is None

    tailored = run.result
    assert tailored.grounding_status == GROUNDING_FAILED
    assert tailored.error_count >= 1
    assert tailored.artifact is None
    assert any(issue.severity == SEVERITY_ERROR for issue in tailored.issues)


def test_a_rejected_document_is_kept_so_the_user_can_see_it(
    db, user, setup, storage, monkeypatch
):
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements, generator="ai")

    class LyingTailor:
        def generate(self, inputs):
            return GenerationResult(
                document=TailoredResumeDocument(
                    experience=[
                        TailoredExperience(
                            source_id=999,
                            company="Somewhere Else",
                            job_title="VP Engineering",
                            start_date=date(2015, 1, 1),
                        )
                    ]
                ),
                generator="ai",
                model="test-model",
                prompt_version="test",
            )

    monkeypatch.setattr(pipeline, "get_tailor", lambda _generator=None: LyingTailor())

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert run.result.payload["experience"][0]["job_title"] == "VP Engineering"
    assert run.result.grounding_status == GROUNDING_FAILED
    assert run.result.artifact is None


# --- failures ----------------------------------------------------------------------


def test_a_provider_failure_ends_the_run_without_leaking_the_provider_text(
    db, user, setup, storage, monkeypatch
):
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements, generator="ai")

    class FailingTailor:
        def generate(self, inputs):
            raise ProviderError(
                "ai_rate_limited",
                "rate limit exceeded for prompt beginning 'Built the payments...'",
            )

    monkeypatch.setattr(pipeline, "get_tailor", lambda _generator=None: FailingTailor())

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert run.status == STATUS_FAILED
    assert run.failed_stage == "generate"
    assert run.error_code == "ai_rate_limited"
    # The user's own career text must not travel out in an error column that
    # gets logged and shipped to an APM.
    assert "Built the payments" not in run.error_message
    assert service.can_retry(run) is True


def test_a_permanent_failure_is_not_offered_a_retry(db, user, setup, storage, monkeypatch):
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements, generator="ai")

    class RefusedTailor:
        def generate(self, inputs):
            raise ProviderError("ai_content_rejected", "declined")

    monkeypatch.setattr(pipeline, "get_tailor", lambda _generator=None: RefusedTailor())

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert run.status == STATUS_FAILED
    assert service.can_retry(run) is False


def test_a_version_with_nothing_in_it_fails_rather_than_producing_an_empty_resume(
    db, user, storage
):
    version = persisted_profile_version(db, user, skills=[("Python", 3.0)])
    job, requirements = a_backend_job(db, user, external_job_id="empty-1")
    run = a_run(db, user, version, job, requirements)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert run.status == STATUS_FAILED
    assert run.error_code == "empty_profile_version"
    assert run.result is None


def test_a_storage_failure_is_transient_and_retryable(db, user, setup, monkeypatch):
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    class BrokenStorage(InMemoryStorage):
        def put(self, key, data, *, content_type):
            from app.services.storage import StorageError

            raise StorageError("the bucket is having a bad minute")

    pipeline.process_run(db, run_id=run.id, storage=BrokenStorage())
    db.refresh(run)

    assert run.status == STATUS_FAILED
    assert run.failed_stage == "render"
    assert run.error_code == "artifact_storage_failed"
    assert service.can_retry(run) is True


def test_a_rendered_file_that_fails_verification_is_render_failed_with_no_artifact(
    db, user, setup, storage, monkeypatch
):
    """M2. The file is read back before it is stored. A file that does not say
    what the approved document says is not handed over as though it did."""
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    # A renderer that writes a real, valid DOCX with a blank paragraph in it:
    # it opens, and it is not the file `render` means to write.
    real_render = pipeline.render

    def with_a_blank_line(document):
        import io

        from docx import Document

        written = Document(io.BytesIO(real_render(document)))
        written.add_paragraph("")
        buffer = io.BytesIO()
        written.save(buffer)
        return buffer.getvalue()

    monkeypatch.setattr(pipeline, "render", with_a_blank_line)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert run.status == STATUS_FAILED
    assert (run.failed_stage, run.error_code) == ("render", "render_failed")
    assert run.result is None
    assert storage._objects == {}


def test_every_file_the_pipeline_stores_passes_verification(db, user, setup, storage):
    from app.services.tailoring.rendering import verify

    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    document = TailoredResumeDocument.model_validate(run.result.payload)
    verify(storage.get(run.result.artifact.storage_key), document)


def test_a_form_feed_in_the_profile_reaches_the_file_as_a_space(db, user, storage):
    """L5, end to end. PDF extraction leaves a form feed in a description; the
    deterministic generator quotes it verbatim, grounding rightly passes it, and
    the renderer used to crash on it — a permanent `render_failed` for a
    document with nothing wrong with it."""
    import io

    from docx import Document

    version = a_tailorable_version(
        db,
        user,
        experiences=[
            {
                "company": "Acme Corp",
                "job_title": "Backend Engineer",
                "description": "Built the payments\x0cservice in Python for the finance team.",
            }
        ],
    )
    job, requirements = a_backend_job(db, user, external_job_id="formfeed-1")
    run = a_run(db, user, version, job, requirements)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert run.status == STATUS_COMPLETED, run.error_code
    assert run.result.grounding_status == GROUNDING_PASSED
    text = "\n".join(
        p.text for p in Document(io.BytesIO(storage.get(run.result.artifact.storage_key))).paragraphs
    )
    assert "Built the payments service in Python for the finance team." in text


@pytest.mark.parametrize("where,stage", [("validate", "validate"), ("render", "render")])
def test_an_unexpected_error_names_the_stage_it_happened_in(
    db, user, setup, storage, monkeypatch, where, stage
):
    """L3. The stage used to be read after the rollback, when the row said
    `queued` again, so every unexpected failure was reported at "generate"."""
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    def boom(*args, **kwargs):
        raise RuntimeError("something nobody modelled")

    if where == "validate":
        monkeypatch.setattr(pipeline.reconciliation, "settle", boom)
    else:
        monkeypatch.setattr(storage, "put", boom)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert (run.status, run.error_code, run.failed_stage) == ("failed", "unexpected_error", stage)


def _foreign(db, user):
    from tests.matching_fixtures import another_user

    stranger = another_user(db)
    return stranger, a_tailorable_version(db, stranger), *a_backend_job(
        db, stranger, external_job_id="theirs-1"
    )


@pytest.mark.parametrize("broken", ["version", "job", "requirements"])
def test_a_run_naming_inputs_that_are_not_its_own_is_refused(
    db, user, setup, storage, broken
):
    """L2. The worker loads inputs by primary key, so the run row is all that
    stands between it and another user's data. A row naming someone else's
    version or job, or a requirements row of a different job, is failed — and
    nothing is generated from it."""
    version, job, requirements = setup
    _, their_version, their_job, their_requirements = _foreign(db, user)
    other_job, other_requirements = a_backend_job(db, user, external_job_id="mine-2")

    inputs = {
        "version": (their_version, job, requirements),
        "job": (version, their_job, their_requirements),
        "requirements": (version, job, other_requirements),
    }[broken]
    run = a_run(db, user, *inputs)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert (run.status, run.error_code, run.failed_stage) == ("failed", "unexpected_error", "generate")
    assert run.result is None
    assert storage._objects == {}


def test_a_run_naming_a_match_for_other_inputs_is_refused(db, user, setup, storage):
    from app.services.matching.service import create_match

    version, job, requirements = setup
    other_job, other_requirements = a_backend_job(db, user, external_job_id="mine-3")
    match, _ = create_match(db, user_id=user.id, job_id=other_job.id, profile_version_id=version.id)
    run = a_run(db, user, version, job, requirements, match_result_id=match.id)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert (run.status, run.error_code) == ("failed", "unexpected_error")
    assert run.result is None


# --- the state machine ----------------------------------------------------------------


def test_a_terminal_run_is_not_reprocessed(db, user, setup, storage):
    """The queue is advisory, so duplicate delivery is ordinary."""
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)
    first_artifact = run.result.artifact.id

    pipeline.process_run(db, run_id=run.id, storage=storage)
    db.refresh(run)

    assert run.result.artifact.id == first_artifact


def test_an_illegal_transition_is_refused(db, user, setup):
    version, job, requirements = setup
    run = a_run(db, user, version, job, requirements)

    with pytest.raises(pipeline.IllegalTransition):
        pipeline.advance(run, STATUS_COMPLETED)


def test_a_missing_run_is_not_an_error(db, storage):
    assert pipeline.process_run(db, run_id=999_999, storage=storage) is None
