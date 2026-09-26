"""Real documents for the resume tests.

These build actual PDF and DOCX bytes rather than checking in binary fixtures:
the tests then exercise the real parsers, and a reader can see exactly what the
document contains. The PDF writer is hand-rolled because the alternative is a
rendering dependency the application itself does not need.
"""

from io import BytesIO


def _escape(text: str) -> bytes:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)").encode("latin-1", "replace")


def make_pdf(pages: list[str]) -> bytes:
    """A minimal, valid, text-bearing PDF — one content stream per page."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    page_ids: list[int] = []
    content_ids: list[int] = []
    for text in pages:
        lines = text.splitlines() or [""]
        stream_parts = [b"BT", b"/F1 12 Tf", b"72 720 Td", b"14 TL"]
        for line in lines:
            stream_parts.append(b"(" + _escape(line) + b") Tj")
            stream_parts.append(b"T*")
        stream_parts.append(b"ET")
        stream = b"\n".join(stream_parts)
        content_ids.append(add(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)))

    pages_id = len(objects) + len(pages) + 1
    for content_id in content_ids:
        page_ids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
                % (pages_id, font_id, content_id)
            )
        )

    kids = b" ".join(b"%d 0 R" % pid for pid in page_ids)
    actual_pages_id = add(b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids)))
    catalog_id = add(b"<< /Type /Catalog /Pages %d 0 R >>" % actual_pages_id)

    out = BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number)
        out.write(body)
        out.write(b"\nendobj\n")

    xref_offset = out.tell()
    out.write(b"xref\n0 %d\n" % (len(objects) + 1))
    out.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root %d 0 R >>\n" % (len(objects) + 1, catalog_id))
    out.write(b"startxref\n%d\n%%%%EOF\n" % xref_offset)
    return out.getvalue()


def make_image_only_pdf(page_count: int = 1) -> bytes:
    """A PDF with pages but no text — what a scan looks like to a text extractor."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    pages_id_placeholder = page_count + 1
    page_ids = []
    for _ in range(page_count):
        page_ids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] /Resources << >> >>"
                % pages_id_placeholder
            )
        )
    kids = b" ".join(b"%d 0 R" % pid for pid in page_ids)
    pages_id = add(b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids)))
    catalog_id = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)

    out = BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))

    xref_offset = out.tell()
    out.write(b"xref\n0 %d\n" % (len(objects) + 1))
    out.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root %d 0 R >>\n" % (len(objects) + 1, catalog_id))
    out.write(b"startxref\n%d\n%%%%EOF\n" % xref_offset)
    return out.getvalue()


def make_docx(
    *,
    headings_and_paragraphs: list[tuple[str, str]] | None = None,
    table: list[list[str]] | None = None,
) -> bytes:
    """A real DOCX with headings, paragraphs and optionally a table."""
    import docx

    document = docx.Document()
    # `or` would treat an explicitly empty list as "use the default", and an
    # empty document is exactly what one of the tests needs.
    if headings_and_paragraphs is None:
        headings_and_paragraphs = [("heading", "EXPERIENCE"), ("paragraph", "Engineer")]
    for kind, text in headings_and_paragraphs:
        if kind == "heading":
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)

    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                grid.cell(row_index, column_index).text = value

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_with_hyperlink(*, text: str, url: str) -> bytes:
    """A DOCX whose paragraph carries a real external hyperlink.

    Built by hand because python-docx has no public API for adding one: the
    relationship and the w:hyperlink element are what the extractor reads, so a
    fixture that faked either would not exercise it.
    """
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement

    document = docx.Document()
    paragraph = document.add_paragraph()
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


RESUME_TEXT_PAGE_ONE = """SUBHASH KUMAR
subhash@example.com | Vijayawada, India

EXPERIENCE
AI Developer Intern, APTRANSCO
May 2024 - August 2024
Built internal tooling in Python and FastAPI.

SKILLS
Python, Django, PostgreSQL, FastAPI"""

RESUME_TEXT_PAGE_TWO = """EDUCATION
B.Tech, Mechanical Engineering
Anil Neerukonda Institute of Technology and Sciences
2020 - 2024

CERTIFICATIONS
AZ-900, Microsoft"""


def make_resume_pdf() -> bytes:
    return make_pdf([RESUME_TEXT_PAGE_ONE, RESUME_TEXT_PAGE_TWO])


def make_resume_docx() -> bytes:
    return make_docx(
        headings_and_paragraphs=[
            ("paragraph", "SUBHASH KUMAR"),
            ("heading", "EXPERIENCE"),
            ("paragraph", "AI Developer Intern, APTRANSCO"),
            ("paragraph", "May 2024 - August 2024"),
            ("heading", "SKILLS"),
            ("paragraph", "Python, Django, PostgreSQL"),
        ],
        table=[["Skill", "Level"], ["Python", "Advanced"]],
    )
