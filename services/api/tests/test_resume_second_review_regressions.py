"""Regressions the second adversarial review of the resume-ingestion fixes found.

Each is a resume shape the first round of fixes made read WORSE than the
pre-session baseline, or left reading wrong, reproduced offline before it was
changed. The rules they pin stay: a heading is still the whole line (so "Project
Manager" is a job title), a page header is still read first, a certification is
still one per line. These pin where each rule was drawn too tightly or too
loosely.
"""

from __future__ import annotations

import pytest

from app.services.resume.extraction import (
    FURNITURE_LINES_KEY,
    ExtractedDocument,
    ExtractedPage,
)
from app.services.resume.mapping import to_snapshot
from app.services.resume.structuring import DeterministicResumeStructurer, _heading


def _document(text: str, **metadata) -> ExtractedDocument:
    return ExtractedDocument(
        pages=[ExtractedPage(page_number=1, text=text, method="text", metadata=metadata)]
    )


def _reading(text: str, **metadata):
    return DeterministicResumeStructurer().structure(_document(text, **metadata))


def _read(text: str, **metadata):
    return _reading(text, **metadata).candidate


# --- headings: the forms the whole-line rule dropped -------------------------------


@pytest.mark.parametrize(
    ("line", "section"),
    [
        # two columns' headings on one baseline: PDFium joins them with ONE space
        ("EXPERIENCE SKILLS", "experience"),
        ("EDUCATION CERTIFICATIONS", "education"),
        ("PROFESSIONAL EXPERIENCE TECHNICAL SKILLS", "experience"),
        # a Word table row: heading cell | empty cell, or | a second heading
        ("EXPERIENCE | ", "experience"),
        ("SKILLS |", "skills"),
        ("EXPERIENCE | SKILLS", "experience"),
        # curated multi-word forms the prefix match used to accept
        ("Academic Profile", "education"),
        ("ACADEMIC PROFILE", "education"),
        ("Academic Record", "education"),
        ("Academic Records", "education"),
        ("Academic Credentials", "education"),
        ("Languages Known", "achievements"),
        ("LANGUAGES KNOWN", "achievements"),
        ("Awards & Recognition", "achievements"),
        ("Awards and Recognitions", "achievements"),
        ("Honors & Recognition", "achievements"),
        ("Publications & Presentations", "achievements"),
        ("Certification Courses", "certifications"),
        ("Summary of Experience", "summary"),
        ("Summary of Skills", "summary"),
        ("References available upon request", "achievements"),
        ("References available upon request.", "achievements"),
        # a section resumed on the next page
        ("Experience (continued)", "experience"),
        ("Experience (Contd.)", "experience"),
        ("EXPERIENCE CONTINUED", "experience"),
        ("Work Experience (cont.)", "experience"),
        ("PROFESSIONAL EXPERIENCE (CONTINUED)", "experience"),
        # decoration typed around the words
        ("EXPERIENCE –", "experience"),
        ("EDUCATION ____", "education"),
        ("EDUCATION " + "_" * 60, "education"),
        ("Projects •", "projects"),
        ("— SKILLS —", "skills"),
        ("=== EDUCATION ===", "education"),
        # what already worked keeps working
        ("EXPERIENCE                AWS", "experience"),
        ("SKILLS\tPython", "skills"),
    ],
)
def test_a_common_or_decorated_heading_opens_its_section(line, section):
    assert _heading(line) == section


@pytest.mark.parametrize(
    "line",
    [
        "Project Manager",
        "Project Manager | Acme Corp | 2020 - Present",
        "Projects I led at Acme",
        "Experience building Django apps",
        "Skills: Python, Go",
        "References: Available upon request",
        # a heading word beside the other column's CONTENT at a narrow gap is
        # the same shape as a job title, and stays content
        "EXPERIENCE AWS Certified",
        # one heading and one non-heading cell is a row of content
        "Senior Engineer | Skills",
        "Kubernetes | PROJECTS",
        # a leading bullet alone is a list item, not ornament
        "• Awards",
        "• Languages:",
    ],
)
def test_the_whole_line_rule_still_keeps_job_titles_and_items_out(line):
    assert _heading(line) is None


def test_two_column_headings_on_one_line_do_not_leave_the_resume_in_the_header():
    """The review's PDF: PDFium joined 'EXPERIENCE' and 'SKILLS' with one space,
    and the draft came out holding a contact record and nothing else."""
    candidate = _read(
        "\n".join(
            [
                "Maria Garcia",
                "maria@example.es | Madrid, Spain",
                "EXPERIENCE SKILLS",
                "Data Analyst | Telefonica | 2019 - Present",
                "- Built dashboards for the retail team.",
                "EDUCATION CERTIFICATIONS",
                "BSc, Statistics | 2013 - 2017",
                "Universidad Complutense de Madrid",
            ]
        )
    )
    assert [role.company for role in candidate.experiences] == ["Telefonica"]
    assert [row.institution for row in candidate.education] == ["Universidad Complutense de Madrid"]


def test_a_word_table_heading_row_with_an_empty_cell_opens_its_section():
    candidate = _read(
        "Sofia Rossi\nEXPERIENCE | \nBackend Engineer, Satispay | 2020 - Present\n"
        "SKILLS | \nGo, PostgreSQL, Kafka"
    )
    assert [(role.job_title, role.company) for role in candidate.experiences] == [
        ("Backend Engineer", "Satispay")
    ]
    assert [skill.name for skill in candidate.skills] == ["Go", "PostgreSQL", "Kafka"]
    # The first experience line is no longer read as the candidate's location.
    assert candidate.contact.location is None


def test_an_academic_profile_after_skills_is_education_not_more_skills():
    reading = _reading(
        "\n".join(
            [
                "Ravi Kumar Reddy",
                "ravi.reddy@gmail.com | Hyderabad, India",
                "TECHNICAL SKILLS",
                "Languages: Python, Java, C",
                "ACADEMIC PROFILE",
                "B.Tech, Computer Science and Engineering | CGPA: 8.1",
                "JNTU Kakinada | 2018 - 2022",
                "LANGUAGES KNOWN",
                "English, Hindi, Telugu",
            ]
        )
    )
    assert {skill.name for skill in reading.candidate.skills} == {"Python", "Java", "C"}
    assert reading.candidate.education[0].degree == "B.Tech"
    # Spoken languages are reported, never filed as skills.
    assert any("English" in item.summary for item in reading.unmapped)


def test_awards_after_experience_are_reported_not_read_as_jobs():
    reading = _reading(
        "Jane Doe\njane@doe.com\nEXPERIENCE\nSoftware Engineer, Infosys | Jun 2019 - Present\n"
        "- Built REST APIs in Django.\nAWARDS & RECOGNITION\nEmployee of the Year, Infosys | 2021\n"
        "Best Paper Award, IEEE ICML | 2020"
    )
    assert [role.job_title for role in reading.candidate.experiences] == ["Software Engineer"]
    assert any("Employee of the Year" in item.summary for item in reading.unmapped)


def test_experience_continued_on_page_two_is_experience_again():
    document = ExtractedDocument(
        pages=[
            ExtractedPage(
                page_number=1,
                text="Jane Doe\nEXPERIENCE\nSenior Engineer, Acme Corp | Jan 2020 - Present\n"
                "- Led the migration.\nSKILLS\nPython, Go, Kubernetes",
                method="text",
            ),
            ExtractedPage(
                page_number=2,
                text="EXPERIENCE (CONTINUED)\nSoftware Engineer, Globex Inc | Jun 2016 - Dec 2019\n"
                "- Built billing services in Go.",
                method="text",
            ),
        ]
    )
    candidate = DeterministicResumeStructurer().structure(document).candidate
    assert [role.company for role in candidate.experiences] == ["Acme Corp", "Globex Inc"]
    assert {skill.name for skill in candidate.skills} == {"Python", "Go", "Kubernetes"}


def test_references_available_upon_request_is_not_a_skill():
    candidate = _read(
        "Jane Doe\nSKILLS\nTools: Git, Docker, Jenkins\nReferences available upon request"
    )
    assert [skill.name for skill in candidate.skills] == ["Git", "Docker", "Jenkins"]


# --- certifications: LinkedIn's layout, and issuers nobody named ------------------------


def _certifications(*lines: str):
    return _read("\n".join(["Jane Doe", "Certifications", *lines])).certifications


def test_a_linkedin_certification_list_reads_as_its_certifications_and_no_more():
    """Name / issuer / "Issued … · Expires …" / "Credential ID …", as LinkedIn
    writes it — with AWS as "Amazon Web Services (AWS)" and a credential that
    never expires as "No Expiration Date"."""
    certifications = _certifications(
        "Microsoft Certified: Azure Fundamentals",
        "Microsoft",
        "Issued Jan 2021 · No Expiration Date",
        "Credential ID H7K2-99",
        "AWS Certified Solutions Architect – Associate",
        "Amazon Web Services (AWS)",
        "Issued Mar 2022 · Expires Mar 2025",
        "Show credential",
    )
    assert [
        (
            item.name,
            item.issuing_organization,
            item.issue_date and item.issue_date.raw,
            item.expiration_date and item.expiration_date.raw,
            item.credential_id,
        )
        for item in certifications
    ] == [
        ("Microsoft Certified: Azure Fundamentals", "Microsoft", "Jan 2021", None, "H7K2-99"),
        (
            "AWS Certified Solutions Architect – Associate",
            "Amazon Web Services (AWS)",
            "Mar 2022",
            "Mar 2025",
            None,
        ),
    ]


def test_a_credential_url_line_is_the_credentials_url_not_a_certification():
    (item,) = _certifications(
        "AWS Certified Developer – Associate",
        "Credential URL: https://www.credly.com/badges/abc",
    )
    assert item.credential_url == "https://www.credly.com/badges/abc"


@pytest.mark.parametrize("line", ["No Expiration Date", "Does not expire", "Valid until: Lifetime", "Expires: Never"])
def test_a_stated_absence_of_expiry_joins_the_certification_above_it(line):
    (item,) = _certifications("Oracle Certified Professional", line)
    assert item.expiration_date is None


@pytest.mark.parametrize(
    "lines",
    [
        ("PMP", "ITIL", "Six Sigma Green Belt"),
        ("Certified ScrumMaster (CSM)", "SAFe"),
        ("Tableau Desktop Specialist", "Tableau"),
        ("CKA", "Kubernetes"),
    ],
)
def test_a_credential_named_like_an_organisation_is_a_certification_not_an_issuer(lines):
    """ITIL, SAFe, Tableau and Kubernetes are listed bare as the credential; as
    known issuers they became the line above's issuer, and vanished."""
    certifications = _certifications(*lines)
    assert [item.name for item in certifications] == list(lines)
    assert all(item.issuing_organization is None for item in certifications)


def test_a_by_line_that_is_a_sentence_is_not_an_issuer():
    (first, second) = _certifications(
        "Google Data Analytics Professional Certificate",
        "by Google on Coursera, completed 2022",
        "Deep Learning Specialization",
    )
    assert first.issuing_organization is None
    assert second.name == "Deep Learning Specialization"


def test_a_by_line_naming_a_known_issuer_and_an_explicit_label_are_still_issuers():
    certifications = _certifications(
        "Machine Learning Specialization",
        "by Coursera",
        "Professional Scrum Master I",
        "Issued by: Scrum.org",
    )
    assert [item.issuing_organization for item in certifications] == ["Coursera", "Scrum.org"]


# --- DOCX: a link in a table cell is where its cell puts it ---------------------------------


def _hyperlink(paragraph, url: str, caption: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    paragraph._p.append(
        parse_xml(
            f'<w:hyperlink {nsdecls("w", "r")} r:id="{rel}"><w:r><w:t>{caption}</w:t></w:r></w:hyperlink>'
        )
    )


def _docx_bytes(document) -> bytes:
    from io import BytesIO

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_a_project_link_in_a_sidebar_table_is_not_the_candidates_own_link():
    """The whole body in one two-column row (sidebar | main): every link in the
    row used to be 'before any heading', so a project's GitHub and Demo became
    the candidate's personal links."""
    import docx

    from app.services.resume.extraction import extract_docx
    from app.services.resume.structuring import unplaced_links

    document = docx.Document()
    document.add_paragraph("Chen Wei")
    _hyperlink(document.add_paragraph("chen@example.com | "), "https://linkedin.com/in/chenwei", "LinkedIn")
    left, right = document.add_table(rows=1, cols=2).rows[0].cells
    left.paragraphs[0].text = "SKILLS"
    left.add_paragraph("Python")
    right.paragraphs[0].text = "PROJECTS"
    _hyperlink(
        right.add_paragraph("Route Planner | Python, Flask "),
        "https://github.com/chenwei/route-planner",
        "GitHub",
    )
    _hyperlink(right.add_paragraph("Chat App | React, Node "), "https://chat-demo.example.com", "Demo")

    extracted = extract_docx(_docx_bytes(document))
    candidate = DeterministicResumeStructurer().structure(extracted).candidate

    assert [link.url for link in candidate.contact.links] == ["https://linkedin.com/in/chenwei"]
    # Not guessed onto anything else either: reported on the review screen.
    reported = " ".join(item.summary for item in unplaced_links(candidate, extracted))
    assert "route-planner" in reported and "chat-demo" in reported


def test_a_contact_link_at_the_top_of_a_table_cell_is_still_a_contact_link():
    import docx

    from app.services.resume.extraction import extract_docx

    document = docx.Document()
    left, right = document.add_table(rows=1, cols=2).rows[0].cells
    left.paragraphs[0].text = "Chen Wei"
    _hyperlink(left.add_paragraph("chen@example.com | "), "https://linkedin.com/in/chenwei", "LinkedIn")
    right.paragraphs[0].text = "EXPERIENCE"
    right.add_paragraph("Engineer, ByteDance | 2019 - Present")

    candidate = DeterministicResumeStructurer().structure(
        extract_docx(_docx_bytes(document))
    ).candidate
    assert [link.url for link in candidate.contact.links] == ["https://linkedin.com/in/chenwei"]


# --- the name: page furniture needs evidence ---------------------------------------------


def _name(text: str, furniture_lines: int) -> str | None:
    contact = _read(text, **{FURNITURE_LINES_KEY: furniture_lines}).contact
    return contact.full_name if contact else None


def test_a_company_name_in_the_page_header_is_not_the_candidates_name():
    assert (
        _name(
            "Randstad Technologies\nCandidate Profile - Confidential\nSummary\n"
            "Data engineer with eight years of platform work.",
            2,
        )
        is None
    )


def test_an_agency_header_with_its_own_contact_details_does_not_outrank_the_body():
    assert (
        _name(
            "Acme Recruitment Ltd\njobs@acme-recruitment.co.uk | 020 7946 0000\n"
            "Jane Doe\njane.doe@example.com\nExperience",
            2,
        )
        == "Jane Doe"
    )


def test_a_header_name_its_email_confirms_is_not_replaced_by_a_body_tagline():
    assert (
        _name(
            "Priya Sharma\npriya@example.com\nSenior Data Scientist\nExperience\n"
            "Data Scientist, Acme | 2020 - Present",
            2,
        )
        == "Priya Sharma"
    )


def test_a_header_name_is_confirmed_through_accents_and_initials():
    assert _name("José García\njgarcia@example.es\nExperience", 2) == "José García"


def test_without_page_furniture_the_first_line_is_still_the_name():
    assert _name("Jane Doe\njane@doe.com\nExperience", 0) == "Jane Doe"


# --- one-line roles in columns --------------------------------------------------------------


def _roles(*lines: str):
    reading = _reading("\n".join(["Jane Doe", "Experience", *lines]))
    return reading.candidate.experiences, reading.unmapped


def test_company_then_title_is_read_the_way_round_the_resume_wrote_it():
    roles, _ = _roles(
        "Infosys | Senior Engineer | Jun 2019 - Present",
        "- Built APIs.",
        "Tata Consultancy Services | Software Engineer | Jul 2016 - May 2019",
    )
    assert [(role.job_title, role.company) for role in roles] == [
        ("Senior Engineer", "Infosys"),
        ("Software Engineer", "Tata Consultancy Services"),
    ]


def test_title_then_company_is_unchanged():
    roles, _ = _roles("Project Manager | Acme Corp | 2020 - Present", "- Led a team.")
    assert [(role.job_title, role.company) for role in roles] == [("Project Manager", "Acme Corp")]


def test_a_third_column_is_placed_or_reported_never_dropped():
    roles, unmapped = _roles(
        "Software Engineer | Infosys | Bangalore",
        "Jun 2021 - Present",
        "- Built APIs.",
        "Data Analyst | Deloitte | Full-time | 2019 - 2021",
        "- Built dashboards.",
        "Backend Engineer | Acme Corp | Austin, TX | Contract",
        "Jan 2018 - Dec 2018",
        "- Built services.",
    )
    assert [(role.company, role.location, role.employment_type) for role in roles] == [
        ("Infosys", None, None),
        ("Deloitte", None, "full_time"),
        ("Acme Corp", "Austin, TX", "contract"),
    ]
    assert [item.summary for item in unmapped] == ["Bangalore (Software Engineer | Infosys)"]


# --- skills: a wrapped line is a line that ran out of room --------------------------------


def _skills(*lines: str, page_width_line: str | None = None):
    body = ["Jane Doe", *([page_width_line] if page_width_line else []), "Technical Skills", *lines]
    return [(skill.category, skill.name) for skill in _read("\n".join(body)).skills]


#: A summary-width line, so the page is as wide as a real resume's.
_WIDE = "Backend engineer with eight years of building Python services and data platforms."


def test_a_sub_heading_after_a_short_list_is_not_glued_onto_it():
    assert _skills(
        "Programming Languages",
        "Python, Java, JavaScript, TypeScript",
        "Tools & Technologies",
        "Docker, Jenkins, Git",
    ) == [
        ("Programming Languages", "Python"),
        ("Programming Languages", "Java"),
        ("Programming Languages", "JavaScript"),
        ("Programming Languages", "TypeScript"),
        ("Tools & Technologies", "Docker"),
        ("Tools & Technologies", "Jenkins"),
        ("Tools & Technologies", "Git"),
    ]


def test_skills_one_per_line_after_a_list_are_not_its_tail():
    assert _skills("Languages: Python, Java, JavaScript, TypeScript", "Git", "Docker") == [
        ("Languages", "Python"),
        ("Languages", "Java"),
        ("Languages", "JavaScript"),
        ("Languages", "TypeScript"),
        (None, "Git"),
        (None, "Docker"),
    ]


def test_a_list_well_short_of_the_page_width_did_not_wrap():
    assert _skills("Languages: Python, Java, JavaScript, TypeScript", "Git", page_width_line=_WIDE)[-1] == (
        None,
        "Git",
    )


def test_a_list_that_ran_to_the_page_edge_still_wraps():
    assert _skills(
        "Cloud & Tooling: AWS, Docker, Kubernetes, Terraform, GitHub Actions, Load", "Testing"
    )[-1] == ("Cloud & Tooling", "Load Testing")


# --- the prompt's fence: neutralised on the way in, restored on the way out ---------------


class _CopyingProvider:
    """A model that obeys rule 3 to the letter: it quotes, EXACTLY, the line it
    was shown."""

    name = "fake"
    model = "fake-model"

    def complete(self, *, system, user, output_model):
        from app.schemas.resume_candidate import CandidateExperience, Provenance
        from app.services.resume.providers import StructuredCompletion

        shown = next(line for line in user.splitlines() if "components in React" in line)
        assert "‹Page" in shown  # the fence is still closed on the way in
        return StructuredCompletion(
            output=output_model(
                experiences=[
                    CandidateExperience(
                        job_title="Frontend Engineer",
                        company="Acme",
                        description=shown,
                        provenance=[Provenance(page=1, quote=shown, confidence=0.9)],
                    )
                ]
            ),
            metadata={},
        )


def test_a_faithful_quote_of_a_neutralised_tag_is_verified_and_stored_as_written():
    from app.services.resume.structuring import ProviderResumeStructurer, verify_provenance

    line = "Built reusable <Page /> and <Layout /> components in React."
    document = _document(f"Jane Doe\nExperience\nFrontend Engineer, Acme\n{line}")
    candidate = verify_provenance(
        ProviderResumeStructurer(_CopyingProvider()).structure(document).candidate, document
    )

    (role,) = candidate.experiences
    assert role.description == line
    assert role.provenance[0].quote == line
    assert role.provenance[0].verified is True


# --- encrypted PDFs: the page tree, and encryption pypdf does not implement ------------------


def _owner_password_pdf(pages: list[str]) -> bytes:
    from io import BytesIO

    from pypdf import PdfReader, PdfWriter

    from tests.resume_fixtures import make_pdf

    writer = PdfWriter(clone_from=PdfReader(BytesIO(make_pdf(pages))))
    writer.encrypt(user_password="", owner_password="owner-secret", algorithm="RC4-128")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _raw_pdf(objects: list[bytes], trailer: bytes) -> bytes:
    from io import BytesIO

    out = BytesIO()
    out.write(b"%PDF-1.7\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R %s >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, trailer, xref)
    )
    return out.getvalue()


@pytest.mark.parametrize(("declared", "actual"), [(1, 3), (3, 1)])
def test_an_owner_password_pdf_is_read_by_its_page_tree_not_its_declared_count(declared, actual):
    """An understated /Count dropped pages silently; an overstated one raised
    IndexError, an unexpected_error retried to exhaustion."""
    from app.services.resume.extraction import extract_pdf

    data = _owner_password_pdf([f"Page {n} " * 20 for n in range(1, actual + 1)])
    assert data.count(b"/Count %d" % actual) == 1
    document = extract_pdf(data.replace(b"/Count %d" % actual, b"/Count %d" % declared, 1))

    assert document.page_count == actual
    assert document.pages_not_read == 0
    assert [page.text.split()[1] for page in document.pages] == [str(n) for n in range(1, actual + 1)]


@pytest.mark.parametrize(
    "encrypt",
    [
        # a certificate security handler: pypdf has only the Standard one
        b"<</Filter/Adobe.PubSec/SubFilter/adbe.pkcs7.s5/V 4/Length 128/Recipients[<3000>]>>",
        # an algorithm version pypdf does not implement
        b"<</Filter/Standard/V 9/R 9/Length 128/P -1028/O<" + b"00" * 32 + b">/U<" + b"00" * 32 + b">>>",
    ],
)
def test_encryption_pypdf_does_not_implement_is_a_permanent_encrypted_failure(encrypt):
    from app.services.resume import failures
    from app.services.resume.extraction import ExtractionError, extract_pdf

    data = _raw_pdf(
        [
            b"<</Type/Catalog/Pages 2 0 R>>",
            b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]>>",
            encrypt,
        ],
        b"/Encrypt 4 0 R /ID[<0011><0011>]",
    )
    with pytest.raises(ExtractionError) as caught:
        extract_pdf(data)
    assert caught.value.code == "document_encrypted"
    assert not failures.is_retryable(caught.value.code)


# --- DOCX: a header that cannot be read never costs the body ---------------------------------


def test_a_dangling_page_header_reference_does_not_fail_the_body():
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from app.services.resume.extraction import FURNITURE_LINES_KEY, extract_docx

    document = docx.Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("jane@example.com")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Engineer, Acme")
    reference = OxmlElement("w:headerReference")
    reference.set(qn("w:type"), "default")
    reference.set(qn("r:id"), "rId999")  # no such relationship
    document.element.body.sectPr.insert(0, reference)

    page = extract_docx(_docx_bytes(document)).pages[0]
    assert page.text.splitlines() == ["Jane Doe", "jane@example.com", "## Experience", "Engineer, Acme"]
    assert page.metadata[FURNITURE_LINES_KEY] == 0


# --- OCR: a page over the bitmap bound is rendered smaller, not skipped -----------------------


class _CountingOcr:
    name = "fake-ocr"

    def __init__(self):
        self.calls = 0

    def is_available(self):
        return True

    def image_to_text(self, image):
        self.calls += 1
        return "Jane Doe Software Engineer " * 5


def test_a_scan_sized_in_pixels_is_rendered_within_the_bound_and_read():
    """An A4 scan at 300 dpi from an image-to-PDF tool is a 2480 x 3508-point
    page: 35 MP at scale 2, over the bound. It was skipped, and the run failed
    with a permanent no_text_extracted on a machine with an OCR engine."""
    from io import BytesIO

    from pypdf import PdfWriter

    from app.services.resume.extraction import UNREAD_KEY, extract_pdf
    from app.services.resume.ocr import apply_ocr

    writer = PdfWriter()
    writer.add_blank_page(2480, 3508)
    buffer = BytesIO()
    writer.write(buffer)
    data = buffer.getvalue()

    engine = _CountingOcr()
    result = apply_ocr(extract_pdf(data), source=data, page_numbers=[1], engine=engine)

    assert engine.calls == 1
    assert result.total_chars > 0
    assert UNREAD_KEY not in result.pages[0].metadata


@pytest.mark.parametrize(("width", "height"), [(612, 792), (2480, 3508), (3024, 4032), (14400, 14400), (14400, 3)])
def test_the_render_scale_never_asks_for_more_than_the_bound(width, height):
    import math

    from app.services.resume.ocr import MAX_RENDER_PIXELS, _render_scale

    scale = _render_scale(width, height, 2.0)
    assert 0 < scale <= 2.0
    assert math.ceil(width * scale) * math.ceil(height * scale) <= MAX_RENDER_PIXELS
    if width * height * 4 <= MAX_RENDER_PIXELS:
        assert scale == 2.0  # an ordinary page keeps the configured scale


@pytest.mark.parametrize(("width", "height"), [(0, 792), (612, -1), (float("inf"), 792)])
def test_a_page_with_no_renderable_size_is_not_rendered(width, height):
    from app.services.resume.ocr import _render_scale

    assert _render_scale(width, height, 2.0) is None


# --- without OCR: a page with no text is unread, not sparse -----------------------------------


def test_a_page_with_no_text_at_all_is_not_the_sparse_page_the_allowance_is_for():
    from app.services.resume.extraction import must_fail_without_ocr

    def pages(*texts: str) -> ExtractedDocument:
        return ExtractedDocument(
            pages=[ExtractedPage(number, text, "text") for number, text in enumerate(texts, start=1)]
        )

    body = "Software Engineer, Initech. Built payment APIs in Python for the billing platform."
    # A wholly unread page 1 of 2: the name, contact block and current role.
    assert must_fail_without_ocr(pages("", body), [1])
    # A sparse page of its own words is still allowed through.
    assert not must_fail_without_ocr(pages(body, "References on request."), [2])
    assert not must_fail_without_ocr(pages(body, body, "x"), [3])


def test_a_resume_whose_first_page_is_an_unreadable_scan_fails_without_an_engine(db, user):
    from app.services.resume import pipeline, service
    from app.services.resume.validation import PDF_MEDIA_TYPE
    from app.services.storage import InMemoryStorage
    from tests.resume_fixtures import make_pdf

    class NoOcr:
        name = "none"

        def is_available(self):
            return False

        def image_to_text(self, image):  # pragma: no cover - never reached
            raise AssertionError("OCR is unavailable")

    storage = InMemoryStorage()
    data = make_pdf(
        ["", "Software Engineer, Initech\nJun 2017 - Dec 2019\nBuilt payment APIs in Python and PostgreSQL."]
    )
    document, _ = service.upload_resume(
        db, user_id=user.id, filename="resume.pdf", declared_media_type=PDF_MEDIA_TYPE, data=data, storage=storage
    )
    run = pipeline.process_run(
        db, run_id=service.start_run(db, document=document).id, storage=storage, ocr_engine=NoOcr()
    )

    assert run.status == "failed"
    assert (run.failed_stage, run.error_code) == ("ocr", "ocr_unavailable")
