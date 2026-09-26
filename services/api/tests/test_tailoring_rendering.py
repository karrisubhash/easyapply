"""The DOCX renderer, read back the way a recipient's tooling reads it.

No database and no generator: each test hands `render` a document built by hand
and opens the bytes with python-docx. Every defect here was invisible to the
suites that only asserted "it starts with PK" — the file opened, it was simply
wrong about something a reader of it would notice.

  * H3 — an education entry dated once printed "(2018 – Present)", a claim of
    current enrolment nobody made; one dated only by its end lost its date.
  * H4 — the link line printed "LinkedIn | GitHub | Portfolio" with no address
    and no hyperlink behind any of it.
  * M3 — every file said python-docx wrote it, in File > Properties.
  * L5 — a form feed left by PDF extraction crashed the render outright.
  * L6 — every file opened with an empty paragraph under the header.
  * M2 — `verify`, the check that now reads every file before it is stored.
"""

from __future__ import annotations

import io
from datetime import date

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn

from app.schemas.tailored_resume import (
    TailoredCertification,
    TailoredContact,
    TailoredEducation,
    TailoredExperience,
    TailoredLine,
    TailoredLink,
    TailoredResumeDocument,
)
from app.services.tailoring import rendering
from app.services.tailoring.rendering import RenderVerificationError, render, verify


def a_line(text: str) -> TailoredLine:
    return TailoredLine(text=text, source_quote=text, source_field="description")


def a_job(**overrides) -> TailoredExperience:
    base = dict(
        source_id=1,
        company="Acme Corp",
        job_title="Backend Engineer",
        start_date=date(2021, 3, 1),
        start_date_precision="month",
        end_date=date(2024, 6, 1),
        end_date_precision="month",
        bullets=[a_line("Built the payments service in Python.")],
    )
    base.update(overrides)
    return TailoredExperience(**base)


def a_document(**overrides) -> TailoredResumeDocument:
    base = dict(
        contact=TailoredContact(
            full_name="Alex Morgan",
            email="alex@example.com",
            links=[
                TailoredLink(url="https://www.linkedin.com/in/alex-morgan/", label="LinkedIn"),
                TailoredLink(url="https://github.com/alex", label="GitHub"),
            ],
        ),
        experience=[a_job()],
    )
    base.update(overrides)
    return TailoredResumeDocument(**base)


def opened(document: TailoredResumeDocument):
    return Document(io.BytesIO(render(document)))


def text_of(document: TailoredResumeDocument) -> str:
    return "\n".join(paragraph.text for paragraph in opened(document).paragraphs)


def a_degree(**dates) -> TailoredEducation:
    return TailoredEducation(
        source_id=7, institution="University of Bristol", degree="BSc", **dates
    )


# --- H3: a single education date is that date, not a current enrolment ----------------


def test_a_degree_dated_only_by_its_start_prints_that_year_and_no_present():
    text = text_of(a_document(education=[a_degree(start_date=date(2018, 1, 1), start_date_precision="year")]))

    assert "University of Bristol — BSc  (2018)" in text
    assert "Present" not in text


def test_a_degree_dated_only_by_its_end_keeps_its_date():
    text = text_of(a_document(education=[a_degree(end_date=date(2018, 6, 1), end_date_precision="month")]))

    assert "(Jun 2018)" in text


def test_a_degree_with_both_dates_is_a_range_at_each_dates_own_precision():
    text = text_of(
        a_document(
            education=[
                a_degree(
                    start_date=date(2015, 1, 1),
                    start_date_precision="year",
                    end_date=date(2018, 6, 1),
                    end_date_precision="month",
                )
            ]
        )
    )

    assert "(2015 – Jun 2018)" in text
    assert "Jan 2015" not in text


def test_a_degree_with_no_dates_prints_no_brackets():
    text = text_of(a_document(education=[a_degree()]))

    assert "University of Bristol — BSc" in text
    assert "(" not in text.split("University of Bristol")[1].splitlines()[0]


def test_a_current_job_still_reads_present():
    """Experience keeps its reading: a job with no end is the one held now."""
    text = text_of(a_document(experience=[a_job(end_date=None, end_date_precision=None)]))

    assert "Mar 2021 – Present" in text


def test_a_certification_prints_one_year_and_never_present():
    text = text_of(
        a_document(
            certifications=[
                TailoredCertification(source_id=3, name="CKA", issue_date=date(2023, 5, 1))
            ]
        )
    )

    assert "CKA  (2023)" in text
    assert "Present" not in text


# --- H4: a contact link is a real hyperlink that prints its address -------------------


def _hyperlinks(document):
    """(visible text, relationship) for every w:hyperlink in the body."""
    part = document.part
    out = []
    for element in part.element.body.iter(qn("w:hyperlink")):
        text = "".join(node.text or "" for node in element.iter(qn("w:t")))
        out.append((text, part.rels[element.get(qn("r:id"))]))
    return out


def test_every_contact_link_is_an_external_hyperlink_to_its_full_url():
    links = _hyperlinks(opened(a_document()))

    assert [text for text, _ in links] == ["linkedin.com/in/alex-morgan", "github.com/alex"]
    for _, relationship in links:
        assert relationship.reltype == RELATIONSHIP_TYPE.HYPERLINK
        assert relationship.is_external
    assert [relationship.target_ref for _, relationship in links] == [
        "https://www.linkedin.com/in/alex-morgan/",
        "https://github.com/alex",
    ]


def test_the_link_line_prints_addresses_not_labels():
    """The label alone is useless on paper and to a parser reading the text."""
    line = next(
        paragraph.text for paragraph in opened(a_document()).paragraphs if "github" in paragraph.text
    )

    assert line == "linkedin.com/in/alex-morgan  |  github.com/alex"
    assert "LinkedIn" not in line and "GitHub" not in line


@pytest.mark.parametrize(
    "url,shown",
    [
        ("https://www.linkedin.com/in/alex/", "linkedin.com/in/alex"),
        ("http://example.com", "example.com"),
        ("HTTPS://WWW.Example.com/x/", "Example.com/x"),
        ("mailto:work@example.com", "work@example.com"),
        ("tel:+441234567890", "+441234567890"),
        ("github.com/alex/", "github.com/alex"),
        ("https://alex.dev/?ref=cv", "alex.dev/?ref=cv"),
    ],
)
def test_the_visible_text_is_the_address_without_scheme_www_or_trailing_slash(url, shown):
    assert rendering.display_url(url) == shown


def test_a_bare_host_opens_over_https_rather_than_as_a_local_path():
    [(text, relationship)] = _hyperlinks(
        opened(a_document(contact=TailoredContact(full_name="Alex", links=[TailoredLink(url="github.com/alex")])))
    )

    assert text == "github.com/alex"
    assert relationship.target_ref == "https://github.com/alex"


@pytest.mark.parametrize("url", ["file:///C:/secrets.txt", "javascript:alert(1)", "not a url"])
def test_a_link_that_must_not_open_is_printed_but_not_made_clickable(url):
    document = opened(a_document(contact=TailoredContact(full_name="Alex", links=[TailoredLink(url=url)])))

    assert _hyperlinks(document) == []
    assert not [r for r in document.part.rels.values() if r.reltype == RELATIONSHIP_TYPE.HYPERLINK]


# --- M3: the file's properties are the candidate's, not the library's -----------------


def test_the_core_properties_name_the_candidate_and_nothing_else():
    properties = opened(a_document()).core_properties

    assert properties.author == "Alex Morgan"
    assert properties.last_modified_by == "Alex Morgan"
    for field in ("title", "subject", "comments", "keywords", "category"):
        assert getattr(properties, field) == "", field


def test_no_property_mentions_python_docx():
    properties = opened(a_document(contact=None)).core_properties

    assert properties.author == ""
    for field in ("author", "last_modified_by", "title", "subject", "comments", "keywords", "category"):
        assert "python-docx" not in (getattr(properties, field) or ""), field


# --- L5: characters XML cannot hold are dropped, not fatal ----------------------------


def test_a_form_feed_from_pdf_extraction_no_longer_crashes_the_render():
    document = a_document(
        experience=[a_job(bullets=[a_line("Built the payments\x0cservice in Python.\x00")])]
    )

    text = text_of(document)

    assert "Built the payments service in Python." in text
    assert "\x0c" not in text and "\x00" not in text


def test_a_line_that_is_nothing_but_control_characters_leaves_no_empty_paragraph():
    document = a_document(experience=[a_job(bullets=[a_line("\x0c\x01"), a_line("A real bullet here.")])])

    paragraphs = [paragraph.text for paragraph in opened(document).paragraphs]

    assert "A real bullet here." in paragraphs
    assert all(text.strip() for text in paragraphs)


# --- L6: no spacer paragraph ------------------------------------------------------------


def test_the_file_has_no_empty_paragraph_and_the_header_carries_the_gap():
    paragraphs = opened(a_document()).paragraphs

    assert all(paragraph.text.strip() for paragraph in paragraphs)
    link_line = next(p for p in paragraphs if "github.com" in p.text)
    assert link_line.paragraph_format.space_after == rendering.HEADER_GAP


# --- M2: the rendered file is read back before it is stored -----------------------------


def test_a_well_formed_file_verifies():
    document = a_document()
    verify(render(document), document)


def test_a_file_that_does_not_open_is_refused():
    with pytest.raises(RenderVerificationError) as raised:
        verify(b"PK\x03\x04 not really a document", a_document())
    assert raised.value.reason == "does_not_open"


def test_a_file_missing_the_candidate_s_name_is_refused():
    written = render(a_document(contact=TailoredContact(full_name="Someone Else")))

    with pytest.raises(RenderVerificationError) as raised:
        verify(written, a_document())
    assert raised.value.reason == "name_missing"


def _with_body_edit(data: bytes, edit) -> bytes:
    document = Document(io.BytesIO(data))
    edit(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_a_file_with_an_empty_paragraph_is_refused():
    document = a_document()
    written = _with_body_edit(render(document), lambda d: d.add_paragraph(""))

    with pytest.raises(RenderVerificationError) as raised:
        verify(written, document)
    assert raised.value.reason == "empty_paragraph"


def test_a_hyperlink_to_an_internal_or_unopenable_target_is_refused():
    document = a_document()

    def point_inward(d):
        # Re-aim a hyperlink at the document's own styles part: a relationship
        # that exists, and is internal.
        internal = next(r.rId for r in d.part.rels.values() if not r.is_external)
        next(d.part.element.body.iter(qn("w:hyperlink"))).set(qn("r:id"), internal)

    with pytest.raises(RenderVerificationError) as raised:
        verify(_with_body_edit(render(document), point_inward), document)
    assert raised.value.reason == "hyperlink_unresolved"

    def point_at_a_file(d):
        d.part.relate_to("file:///C:/x", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    with pytest.raises(RenderVerificationError):
        verify(_with_body_edit(render(document), point_at_a_file), document)
